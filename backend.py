# backend.py – VokabelMeister Backend (Datenmodell, Lernalgorithmus, JSON-Speicher, Datei-Import)
from __future__ import annotations

import json
import random
import re
from dataclasses import dataclass
from pathlib import Path

# Speicherort für Benutzerdaten (Store-kompatibel, funktioniert auf jedem PC)
import sys as _sys
import os as _os
import shutil as _shutil

def _get_app_dir() -> Path:
    """Ermittelt das Datenverzeichnis:
    1. Liegt neben der .exe / dem Skript eine 'vokabeln.json', wird dieser Ordner genutzt.
    2. Liegt in %LOCALAPPDATA%\\VokabelMeister eine 'vokabeln.json', wird diese genutzt.
    3. Andernfalls der Ordner der .exe (falls beschreibbar) oder %LOCALAPPDATA%\\VokabelMeister.
    """
    if getattr(_sys, "frozen", False):
        exe_dir = Path(_sys.executable).parent
    else:
        exe_dir = Path(__file__).parent

    if (exe_dir / "vokabeln.json").exists():
        return exe_dir

    local_app = Path(_os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local")) / "VokabelMeister"
    if (local_app / "vokabeln.json").exists():
        return local_app

    try:
        test = exe_dir / ".write_test"
        test.touch()
        test.unlink()
        return exe_dir
    except Exception:
        local_app.mkdir(parents=True, exist_ok=True)
        return local_app

APP_DIR = _get_app_dir()
DATA_FILE = APP_DIR / "vokabeln.json"
KATEGORIEN_FILE = APP_DIR / "kategorien.json"


# Häufig genutzte Sprachen
STANDARD_SPRACHEN: list[str] = [
    "Englisch", "Deutsch", "Französisch", "Spanisch", "Italienisch",
    "Latein", "Russisch", "Japanisch", "Chinesisch", "Türkisch",
    "Portugiesisch", "Niederländisch", "Griechisch", "Polnisch", "Arabisch",
]

# Standard-Kategorien mit Fragetext für Vorder- und Rückseite
DEFAULT_KATEGORIEN: dict[str, dict[str, str]] = {
    "Sprache": {
        "icon":         "🌍",
        "beschreibung": "Fremdwörter und Übersetzungen",
        "frage_front":  "Wie lautet die Übersetzung von:",
        "frage_back":   "In welcher Sprache heißt der Begriff:",
        "lbl_front":    "Fremdsprache",
        "lbl_back":     "Zielsprache",
    },
    "Fachwörter": {
        "icon":         "📚",
        "beschreibung": "Fachbegriffe und ihre Bedeutung",
        "frage_front":  "Was bedeutet der Begriff:",
        "frage_back":   "Welcher Begriff beschreibt folgendes:",
        "lbl_front":    "Fachbegriff",
        "lbl_back":     "Bedeutung",
    },
    "Formeln": {
        "icon":         "🔬",
        "beschreibung": "Formeln und ihre Anwendung",
        "frage_front":  "Erkläre die Formel:",
        "frage_back":   "Welche Formel beschreibt folgendes:",
        "lbl_front":    "Formel",
        "lbl_back":     "Beschreibung",
    },
    "Befehle": {
        "icon":         "💻",
        "beschreibung": "Tastenkürzel und was sie machen",
        "frage_front":  "Was macht das Tastenkürzel:",
        "frage_back":   "Welches Tastenkürzel macht folgendes:",
        "lbl_front":    "Tastenkürzel",
        "lbl_back":     "Was sie machen",
    },
}

def load_kategorien() -> dict[str, dict[str, str]]:
    """Lädt Standard- und benutzerdefinierte Kategorien."""
    APP_DIR.mkdir(parents=True, exist_ok=True)
    kats = {k: dict(v) for k, v in DEFAULT_KATEGORIEN.items()}
    if KATEGORIEN_FILE.exists():
        try:
            with open(KATEGORIEN_FILE, "r", encoding="utf-8") as f:
                saved = json.load(f)
            if isinstance(saved, dict):
                kats.update(saved)
        except Exception:
            pass
    return kats

def save_kategorien(kategorien: dict[str, dict[str, str]]) -> None:
    """Speichert benutzerdefinierte Kategorien in JSON."""
    APP_DIR.mkdir(parents=True, exist_ok=True)
    custom = {k: v for k, v in kategorien.items() if k not in DEFAULT_KATEGORIEN or v != DEFAULT_KATEGORIEN[k]}
    with open(KATEGORIEN_FILE, "w", encoding="utf-8") as f:
        json.dump(custom, f, ensure_ascii=False, indent=2)

KATEGORIEN: dict[str, dict[str, str]] = load_kategorien()
KATEGORIE_NAMEN: list[str] = list(KATEGORIEN.keys())

# Standard-Beispielvokabeln für den ersten Start
DEFAULT_VOKABELN: list[dict] = [
    {"front": "apple",                    "back": "Apfel",                              "kategorie": "Sprache",    "lang_front": "Englisch", "lang_back": "Deutsch", "attempts": 0, "correct": 0},
    {"front": "house",                    "back": "Haus",                               "kategorie": "Sprache",    "lang_front": "Englisch", "lang_back": "Deutsch", "attempts": 0, "correct": 0},
    {"front": "Photosynthese",            "back": "Pflanzen erzeugen mit Licht Zucker", "kategorie": "Fachwörter", "lang_front": "",         "lang_back": "",        "attempts": 0, "correct": 0},
    {"front": "E = mc²",                  "back": "Masse-Energie-Äquivalenz (Einstein)","kategorie": "Formeln",    "lang_front": "",         "lang_back": "",        "attempts": 0, "correct": 0},
    {"front": "a² + b² = c²",             "back": "Satz des Pythagoras",                "kategorie": "Formeln",    "lang_front": "",         "lang_back": "",        "attempts": 0, "correct": 0},
    {"front": "Windows + D",              "back": "Desktop anzeigen",                   "kategorie": "Befehle",    "lang_front": "",         "lang_back": "",        "attempts": 0, "correct": 0},
    {"front": "Insert + F7",              "back": "Linkliste öffnen in JAWS",           "kategorie": "Befehle",    "lang_front": "",         "lang_back": "",        "attempts": 0, "correct": 0},
    {"front": "Insert + Pfeil nach unten","back": "Alles vorlesen ab Cursor in JAWS",   "kategorie": "Befehle",    "lang_front": "",         "lang_back": "",        "attempts": 0, "correct": 0},
]


@dataclass
class Vokabel:
    front:      str
    back:       str
    kategorie:  str = "Sprache"
    lang_front: str = "Englisch"  # Sprache der Vorderseite (linke Sprache)
    lang_back:  str = "Deutsch"   # Sprache der Rückseite (rechte Sprache)
    attempts:   int = 0           # Anzahl Abfragen
    correct:    int = 0           # Anzahl richtige Antworten

    @property
    def score(self) -> float:
        # Prozentwert 0–100, bei neuen Vokabeln 0
        if self.attempts == 0:
            return 0.0
        return round((self.correct / self.attempts) * 100.0, 1)

    @property
    def weight(self) -> float:
        # Gewicht für Zufallsauswahl: schlechte Vokabeln erhalten höheres Gewicht
        return (100.0 - self.score) + 10.0

    def to_dict(self) -> dict:
        # Serialisierung in Dictionary für JSON-Speicherung
        return {
            "front":      self.front,
            "back":       self.back,
            "kategorie":  self.kategorie,
            "lang_front": self.lang_front,
            "lang_back":  self.lang_back,
            "attempts":   self.attempts,
            "correct":    self.correct,
        }

    @classmethod
    def from_dict(cls, d: dict) -> "Vokabel":
        # Deserialisierung aus Dictionary mit Abwärtskompatibilität
        return cls(
            front=d.get("front", ""),
            back=d.get("back", ""),
            kategorie=d.get("kategorie", "Sprache"),
            lang_front=d.get("lang_front", "Englisch"),
            lang_back=d.get("lang_back", "Deutsch"),
            attempts=int(d.get("attempts", 0)),
            correct=int(d.get("correct", 0)),
        )


def load_vokabeln() -> list[Vokabel]:
    """Lädt Vokabeln aus der JSON-Datei.
    Existiert keine Datei oder ist sie leer, wird eine leere Liste zurückgegeben.
    Es werden niemals automatisch Standardvokabeln injiziert oder hinzugefügt.
    """
    APP_DIR.mkdir(parents=True, exist_ok=True)
    if not DATA_FILE.exists():
        return []
    try:
        with open(DATA_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, list):
            return []
        return [Vokabel.from_dict(d) for d in data if isinstance(d, dict)]
    except Exception:
        return []


def save_vokabeln(vokabeln: list[Vokabel]) -> None:
    # Speichert alle Vokabeln in die JSON-Datei
    APP_DIR.mkdir(parents=True, exist_ok=True)
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump([v.to_dict() for v in vokabeln], f, ensure_ascii=False, indent=2)


def split_line(line: str) -> tuple[str, str] | None:
    # Zeile in (Vorderseite, Rückseite) aufteilen – trennt bei |, Tab, ;, Gedankenstrich oder Minus (egal ob mit/ohne Leerzeichen)
    line = line.strip()
    if not line or line.startswith("#"):
        return None
    # 1. Spezifische Trenner (Pipe, Tab, Semikolon)
    for sep in ["|", "\t", ";"]:
        if sep in line:
            a, b = line.split(sep, 1)
            a, b = a.strip(), b.strip()
            if a and b:
                return a, b
    # 2. Gedankenstriche (En-Dash – und Em-Dash — aus Word)
    for sep in ["\u2013", "\u2014"]:
        if sep in line:
            a, b = line.split(sep, 1)
            a, b = a.strip(), b.strip()
            if a and b:
                return a, b
    # 3. Normales Minus mit Leerzeichen ' - '
    if " - " in line:
        a, b = line.split(" - ", 1)
        a, b = a.strip(), b.strip()
        if a and b:
            return a, b
    # 4. Normales Minus ohne Leerzeichen '-'
    if "-" in line:
        a, b = line.split("-", 1)
        a, b = a.strip(), b.strip()
        if a and b:
            return a, b
    return None


def parse_datei(path: Path) -> list[tuple[str, str]]:
    # Liest Vokabelpaare aus Word- oder Textdatei
    paare: list[tuple[str, str]] = []
    if path.suffix.lower() == ".docx":
        from docx import Document  # type: ignore[import-untyped]
        doc = Document(str(path))
        # 1. Tabellen mit mindestens 2 Spalten
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if len(cells) >= 2 and cells[0] and cells[1]:
                    # Überschriftenzeilen wie "Vorderseite | Rückseite" ignorieren
                    if cells[0].lower() in ["vorderseite", "wort", "frage", "begriff"]:
                        continue
                    paare.append((cells[0], cells[1]))
        # 2. Falls keine Tabellen vorhanden sind, Absätze durchsuchen
        if not paare:
            for para in doc.paragraphs:
                result = split_line(para.text)
                if result:
                    paare.append(result)
    else:
        # Textdatei: erst UTF-8 probieren, Fallback auf Windows-1252
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = path.read_text(encoding="cp1252", errors="replace")
        for line in text.splitlines():
            result = split_line(line)
            if result:
                paare.append(result)
    return paare


class VokabelTrainer:
    def __init__(self) -> None:
        # Kategorien und Vokabeln automatisch aus JSON laden
        self.kategorien: dict[str, dict[str, str]] = load_kategorien()
        self.vokabeln: list[Vokabel] = load_vokabeln()

    def add_kategorie(
        self,
        name: str,
        icon: str,
        beschreibung: str,
        frage_front: str,
        frage_back: str,
        lbl_front: str = "Vorderseite",
        lbl_back: str = "Rückseite",
    ) -> bool:
        """Legt eine neue benutzerdefinierte Kategorie an."""
        clean_name = name.strip()
        if not clean_name or clean_name in self.kategorien:
            return False
        self.kategorien[clean_name] = {
            "icon":         icon.strip() or "🏷️",
            "beschreibung": beschreibung.strip() or "Eigene Kategorie",
            "frage_front":  frage_front.strip() or "Was bedeutet:",
            "frage_back":   frage_back.strip() or "Welcher Begriff beschreibt folgendes:",
            "lbl_front":    lbl_front.strip() or "Vorderseite",
            "lbl_back":     lbl_back.strip() or "Rückseite",
        }
        # Globale Mappings synchronisieren
        KATEGORIEN.clear()
        KATEGORIEN.update(self.kategorien)
        KATEGORIE_NAMEN.clear()
        KATEGORIE_NAMEN.extend(self.kategorien.keys())
        save_kategorien(self.kategorien)
        return True

    def delete_kategorie(self, name: str) -> bool:
        """Löscht eine benutzerdefinierte Kategorie und alle darin enthaltenen Vokabeln."""
        if name not in self.kategorien:
            return False
        del self.kategorien[name]
        self.vokabeln = [v for v in self.vokabeln if v.kategorie != name]
        save_vokabeln(self.vokabeln)
        KATEGORIEN.clear()
        KATEGORIEN.update(self.kategorien)
        KATEGORIE_NAMEN.clear()
        KATEGORIE_NAMEN.extend(self.kategorien.keys())
        save_kategorien(self.kategorien)
        return True

    def get_custom_kategorien(self) -> dict[str, dict[str, str]]:
        """Gibt alle benutzerdefinierten (löschbaren) Kategorien zurück."""
        return {k: v for k, v in self.kategorien.items() if k not in DEFAULT_KATEGORIEN}

    def get_kategorien(self) -> dict[str, dict[str, str]]:
        return dict(self.kategorien)

    def get_kategorie_namen(self) -> list[str]:
        return list(self.kategorien.keys())

    def get_kategorie_info(self, name: str) -> dict[str, str]:
        if name in self.kategorien:
            return self.kategorien[name]
        return DEFAULT_KATEGORIEN.get(name, {
            "icon": "🏷️",
            "beschreibung": "",
            "frage_front": "Was bedeutet:",
            "frage_back": "Welcher Begriff beschreibt folgendes:",
            "lbl_front": "Vorderseite",
            "lbl_back": "Rückseite",
        })

    def is_duplicate(self, front: str, kategorie: str, lang_front: str = "Englisch") -> bool:
        # Prüft ob eine Vokabel mit gleicher Vorderseite in dieser Kategorie (und Sprache) schon existiert
        target = front.strip().lower()
        for v in self.vokabeln:
            if v.kategorie == kategorie and v.front.lower() == target:
                if kategorie != "Sprache" or v.lang_front.lower() == lang_front.strip().lower():
                    return True
        return False

    def get_reversed_lang_pair(self, lang_front: str, lang_back: str) -> tuple[str, str] | None:
        """Gibt die existierende Richtung zurück, wenn die umgekehrte Kombination schon genutzt wird."""
        lf = lang_front.strip().lower()
        lb = lang_back.strip().lower()
        for v in self.vokabeln:
            if v.kategorie == "Sprache":
                vf = v.lang_front.strip().lower()
                vb = v.lang_back.strip().lower()
                if vf == lb and vb == lf:
                    return (v.lang_front.strip(), v.lang_back.strip())
        return None

    def get_most_used_lang_pair(self) -> tuple[str, str]:
        """Gibt die am häufigsten genutzte Sprachkombi (lang_front, lang_back) zurück.
        Fallback: ('Englisch', 'Deutsch')."""
        counts: dict[tuple[str, str], int] = {}
        for v in self.vokabeln:
            if v.kategorie == "Sprache" and v.lang_front and v.lang_back:
                key = (v.lang_front.strip(), v.lang_back.strip())
                counts[key] = counts.get(key, 0) + 1
        if not counts:
            return ("Englisch", "Deutsch")
        return max(counts.items(), key=lambda item: item[1])[0]


    def add_vokabel(
        self,
        front: str,
        back: str,
        kategorie: str = "Sprache",
        lang_front: str = "Englisch",
        lang_back: str = "Deutsch",
    ) -> Vokabel | None:
        # Fügt neue Vokabel hinzu – falls Duplikat, wird sie ignoriert und None zurückgegeben
        clean_front = front.strip()
        clean_back = back.strip()
        if not clean_front or not clean_back:
            return None
        if self.is_duplicate(clean_front, kategorie, lang_front):
            return None
        vok = Vokabel(
            front=clean_front,
            back=clean_back,
            kategorie=kategorie,
            lang_front=lang_front.strip() or "Englisch",
            lang_back=lang_back.strip() or "Deutsch",
        )
        self.vokabeln.append(vok)
        self.save()
        return vok

    def delete_vokabel(self, vokabel: Vokabel) -> None:
        # Einzelne Vokabel löschen
        if vokabel in self.vokabeln:
            self.vokabeln.remove(vokabel)
            self.save()

    def delete_vokabeln(self, to_delete: list[Vokabel]) -> None:
        # Mehrere Vokabeln gleichzeitig löschen
        for v in to_delete:
            if v in self.vokabeln:
                self.vokabeln.remove(v)
        self.save()

    def delete_all(self, kategorie: str | None = None) -> None:
        # Alle Vokabeln (oder alle einer Kategorie) löschen
        if kategorie:
            self.vokabeln = [v for v in self.vokabeln if v.kategorie != kategorie]
        else:
            self.vokabeln = []
        self.save()

    def save(self) -> None:
        # Speichert den aktuellen Stand in die JSON-Datei
        save_vokabeln(self.vokabeln)

    def get_pool(self, kategorie: list[str] | set[str] | str | None = None) -> list[Vokabel]:
        # Vokabeln zurückgeben, optional nach Kategorie(n) gefiltert (unterstützt Einzel- und Mehrfachauswahl)
        if kategorie is None:
            return list(self.vokabeln)
        if isinstance(kategorie, str):
            return [v for v in self.vokabeln if v.kategorie == kategorie]
        kat_set = set(kategorie)
        return [v for v in self.vokabeln if v.kategorie in kat_set]

    def get_next_card(
        self,
        kategorie: list[str] | set[str] | str | None = None,
        exclude: Vokabel | None = None,
    ) -> Vokabel:
        # Gewichtete Zufallsauswahl – schwache Vokabeln häufiger
        # exclude: diese Vokabel nicht nehmen (max-2-Streak-Schutz), sofern Pool > 1
        pool = self.get_pool(kategorie)
        if not pool:
            msg = "in den ausgewählten Kategorien" if isinstance(kategorie, (list, set)) else (f"in {kategorie}" if kategorie else "")
            raise ValueError(f"Keine Vokabeln {msg} vorhanden.".strip())
        if exclude is not None:
            filtered = [v for v in pool if v is not exclude]
            if filtered:
                pool = filtered
        weights = [v.weight for v in pool]
        return random.choices(pool, weights=weights, k=1)[0]

    def record_result(self, vokabel: Vokabel, is_correct: bool) -> None:
        # Ergebnis eintragen und speichern
        vokabel.attempts += 1
        if is_correct:
            vokabel.correct += 1
        self.save()


def import_datei(
    trainer: VokabelTrainer,
    path: Path,
    kategorie: str,
    lang_front: str = "Englisch",
    lang_back: str = "Deutsch",
) -> tuple[int, int]:
    # Importiert Datei in Trainer; gibt (anzahl_neu, anzahl_duplikate) zurück
    paare = parse_datei(path)
    neu_count = 0
    dup_count = 0
    for front, back in paare:
        res = trainer.add_vokabel(front, back, kategorie, lang_front, lang_back)
        if res is not None:
            neu_count += 1
        else:
            dup_count += 1
    return neu_count, dup_count


# ── Intelligente Antworterkennung & Fehlertoleranz ────────────────────────────
STOPWORDS: set[str] = {
    "der", "die", "das", "ein", "eine", "einer", "eines", "einem", "einen",
    "the", "a", "an", "to", "of", "in", "on", "at", "by", "for", "with", "mit",
    "und", "and", "or", "oder", "sich", "zu", "von", "aus", "den", "dem", "des",
}


def damerau_levenshtein(s1: str, s2: str) -> int:
    """Berechnet die Damerau-Levenshtein-Distanz (Einfügen, Löschen, Ersetzen, Dreher)."""
    d: dict[tuple[int, int], int] = {}
    len1, len2 = len(s1), len(s2)
    for i in range(-1, len1 + 1):
        d[(i, -1)] = i + 1
    for j in range(-1, len2 + 1):
        d[(-1, j)] = j + 1
    for i in range(len1):
        for j in range(len2):
            cost = 0 if s1[i] == s2[j] else 1
            d[(i, j)] = min(
                d[(i - 1, j)] + 1,
                d[(i, j - 1)] + 1,
                d[(i - 1, j - 1)] + cost,
            )
            if i > 0 and j > 0 and s1[i] == s2[j - 1] and s1[i - 1] == s2[j]:
                d[(i, j)] = min(d[(i, j)], d[(i - 2, j - 2)] + 1)
    return d[(len1 - 1, len2 - 1)]


def max_allowed_typos(word_len: int) -> int:
    """Erlaubte Tippfehler: 0 bei <=3 Zeichen, 1 bei 4-7 Zeichen, 2 bei >=8 Zeichen."""
    if word_len <= 3:
        return 0
    if word_len <= 7:
        return 1
    return 2


def clean_matching_str(s: str) -> str:
    """Bereinigt Klammern, Satzzeichen und vereinheitlicht Leerzeichen."""
    s = re.sub(r"[\(\)\[\]\{\}<>„“\"\'`´]", "", s)
    s = re.sub(r"[.,!?;:]", " ", s)
    s = re.sub(r"\s*([+=])\s*", r" \1 ", s)
    return " ".join(s.lower().split())


def _expand_brackets(text: str) -> set[str]:
    """Generiert alle Kombinationen von Klammer-Inhalten (mit Inhalt / ohne Inhalt)."""
    pattern = re.compile(r"[\(\[\{]([^\)\]\}]*)[\)\]\}]")

    def _recurse(curr: str) -> set[str]:
        m = pattern.search(curr)
        if not m:
            return {curr}
        content = m.group(1)
        start, end = m.span()
        opt_a = curr[:start] + content + curr[end:]
        opt_b = curr[:start] + curr[end:]
        res = _recurse(opt_a) | _recurse(opt_b)
        if "-" in content:
            opt_c = curr[:start] + content.replace("-", "") + curr[end:]
            res |= _recurse(opt_c)
        return res

    raw_res = _recurse(text)
    out: set[str] = set()
    for r in raw_res:
        cleaned = " ".join(r.split())
        if cleaned:
            out.add(cleaned)
    return out


def _expand_slash_forms(text: str) -> set[str]:
    """Erweitert Formen mit Schrägstrichen (z. B. thai(länder/in), ein/e, Erwachsene/r, der/die/das)."""
    forms = {text}
    m_in = re.search(r"(\w+)/in\b", text, re.IGNORECASE)
    if m_in:
        base = m_in.group(1)
        forms.add(text[:m_in.start()] + base + text[m_in.end():])
        forms.add(text[:m_in.start()] + base + "in" + text[m_in.end():])

    m_e = re.search(r"(\w+)/e\b", text, re.IGNORECASE)
    if m_e:
        base = m_e.group(1)
        forms.add(text[:m_e.start()] + base + text[m_e.end():])
        forms.add(text[:m_e.start()] + base + "e" + text[m_e.end():])

    m_rs = re.search(r"(\w+)/r/s\b", text, re.IGNORECASE)
    if m_rs:
        base = m_rs.group(1)
        forms.add(text[:m_rs.start()] + base + text[m_rs.end():])
        forms.add(text[:m_rs.start()] + base + "r" + text[m_rs.end():])
        forms.add(text[:m_rs.start()] + base + "s" + text[m_rs.end():])

    m_r_s = re.search(r"(\w+)r/s\b", text, re.IGNORECASE)
    if m_r_s:
        stem = m_r_s.group(1)
        forms.add(text[:m_r_s.start()] + stem + text[m_r_s.end():])
        forms.add(text[:m_r_s.start()] + stem + "r" + text[m_r_s.end():])
        forms.add(text[:m_r_s.start()] + stem + "s" + text[m_r_s.end():])

    m_r = re.search(r"(\w+)/r\b", text, re.IGNORECASE)
    if m_r:
        base = m_r.group(1)
        forms.add(text[:m_r.start()] + base + text[m_r.end():])
        forms.add(text[:m_r.start()] + base + "r" + text[m_r.end():])

    if "/" in text and not any([m_in, m_e, m_rs, m_r_s, m_r]):
        forms.add(text.replace("/", " "))
        slash_words = text.split("/")
        for w in slash_words:
            w_str = w.strip()
            if len(w_str) >= 2 or w_str in {"a", "i", "o", "u"}:
                forms.add(w_str)

    return {f.strip() for f in forms if f.strip()}


def get_answer_candidates(expected: str) -> tuple[list[str], list[str], list[str]]:
    """Erzeugt Kandidaten-Vollformen, Synonym-Teile und Einzelwörter aus der Musterlösung."""
    raw = expected.strip().lower()
    full_forms: set[str] = set()
    parts: set[str] = set()
    all_words: set[str] = set()

    bracket_variants = _expand_brackets(raw)

    for bv in bracket_variants:
        for sf in _expand_slash_forms(bv):
            c = clean_matching_str(sf)
            if c:
                full_forms.add(c)
                parts.add(c)
                for w in c.split():
                    all_words.add(w)

        sub_items = re.split(r"[,;]+|\s+-\s+|\boder\b|\bor\b", bv)
        for item in sub_items:
            for sf in _expand_slash_forms(item):
                c = clean_matching_str(sf)
                if c:
                    parts.add(c)
                    for w in c.split():
                        all_words.add(w)
            if "/" in item:
                for slash_p in item.split("/"):
                    c = clean_matching_str(slash_p)
                    if c:
                        parts.add(c)
                        for w in c.split():
                            all_words.add(w)

    return list(full_forms), list(parts), list(all_words)


def check_answer(user_input: str, expected: str) -> tuple[bool, bool]:
    """Prüft eine Benutzereingabe gegen die erwartete Antwort.
    Gibt (ist_richtig, ist_tippfehler_oder_teilantwort) zurück.
    """
    u_clean = clean_matching_str(user_input)
    if not u_clean:
        return False, False

    full_forms, parts, all_words = get_answer_candidates(expected)
    u_words = u_clean.split()

    # 1. Exakte Übereinstimmung mit Vollform oder Synonym-Teil
    if u_clean in full_forms or u_clean in parts:
        return True, False

    # 2. Ein einzelnes Wort von mehreren reicht (außer reine Stoppwörter)
    if len(u_words) == 1 and u_words[0] in all_words:
        if u_words[0] not in STOPWORDS:
            return True, False

    # 3. Beliebige Wortreihenfolge und Teilmengen (z. B. stellen legen tun)
    if len(u_words) >= 1:
        for p in list(full_forms) + list(parts):
            p_words = p.split()
            if len(p_words) > 1:
                if set(u_words) == set(p_words):
                    return True, False
                if len(u_words) > 1 and set(u_words).issubset(set(p_words)):
                    if any(w not in STOPWORDS for w in u_words):
                        return True, False

    # 4. Tippfehler-Erkennung (Fuzzy Matching)
    # 4a. Auf Vollformen oder Synonym-Teilen
    for target in list(full_forms) + list(parts):
        if not target:
            continue
        limit = max_allowed_typos(len(target))
        if limit > 0 and damerau_levenshtein(u_clean, target) <= limit:
            return True, True

    # 4b. Auf Einzelwörtern aus der Musterlösung
    if len(u_words) == 1:
        w = u_words[0]
        for tw in all_words:
            if tw in STOPWORDS:
                continue
            limit = max_allowed_typos(len(tw))
            if limit > 0 and damerau_levenshtein(w, tw) <= limit:
                return True, True

    # 4c. Mehrere Wörter mit Tippfehlern in beliebiger Reihenfolge
    if len(u_words) > 1:
        for p in list(full_forms) + list(parts):
            p_words = p.split()
            if len(p_words) == len(u_words):
                used_indices: set[int] = set()
                matched_all = True
                for uw in u_words:
                    found = False
                    for idx, pw in enumerate(p_words):
                        if idx in used_indices:
                            continue
                        lim = max_allowed_typos(len(pw))
                        if damerau_levenshtein(uw, pw) <= lim:
                            used_indices.add(idx)
                            found = True
                            break
                    if not found:
                        matched_all = False
                        break
                if matched_all:
                    return True, True

    return False, False
